#!/usr/bin/env python3
"""Collect a finished run: validity, render, deterministic checks, structural diff vs original and ground truth."""
import json, sys, hashlib, subprocess, zipfile, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
WORK=Path('/home/user/bench/work/runs'); REF=Path('/home/user/bench/work/ref'); REPO=Path('/home/user/bench/PPTArena'); OUT=Path('/home/user/bench/work/collected')
manifest=json.load(open(WORK/'manifest.json'))

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def render(pptx, outdir):
    outdir=Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    for f in outdir.glob('*'): f.unlink()
    r=subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(outdir),str(pptx)],capture_output=True,text=True,timeout=600)
    pdf=next(outdir.glob('*.pdf'),None)
    if not pdf: return [], (r.stderr or r.stdout)[-300:]
    subprocess.run(['pdftoppm','-r','60','-png',str(pdf),str(outdir/'slide')],capture_output=True,timeout=600)
    return sorted(str(p) for p in outdir.glob('*.png')), ''

def shape_summary(sh):
    d={'name':sh.name,'type':str(sh.shape_type).split('.')[-1].split(' ')[0] if sh.shape_type else 'None',
       'left':round((sh.left or 0)/914400,2),'top':round((sh.top or 0)/914400,2),'w':round((sh.width or 0)/914400,2),'h':round((sh.height or 0)/914400,2)}
    if sh.has_text_frame:
        runs=[(r.text, r.font.size.pt if r.font.size else None, r.font.bold, r.font.name) for p in sh.text_frame.paragraphs for r in p.runs]
        d['text']=sh.text_frame.text[:300]; d['runs']=len(runs)
        sizes={r[1] for r in runs if r[1]}; bolds={r[2] for r in runs}; fonts={r[3] for r in runs if r[3]}
        if sizes: d['font_sizes']=sorted(sizes)
        if True in bolds: d['bold_runs']=sum(1 for r in runs if r[2])
        if fonts: d['fonts']=sorted(fonts)
        try:
            d['align']=sorted({str(p.alignment).split('.')[-1].split(' ')[0] for p in sh.text_frame.paragraphs if p.alignment is not None})
        except Exception: pass
    if getattr(sh,'has_chart',False) and sh.has_chart:
        ch=sh.chart
        try:
            d['chart']={'type':str(ch.chart_type).split('.')[-1].split(' ')[0],'has_legend':ch.has_legend,
                        'series':[{'name':s.name,'n':len(list(s.values)),'values':[round(v,3) if isinstance(v,(int,float)) else v for v in list(s.values)[:12]]} for p in ch.plots for s in p.series],
                        'categories':[str(c) for c in list(ch.plots[0].categories)[:12]] if len(ch.plots) else []}
        except Exception as e: d['chart']={'error':str(e)}
    if getattr(sh,'has_table',False) and sh.has_table:
        t=sh.table; d['table']={'rows':len(t.rows),'cols':len(t.columns),'cells':[[c.text[:40] for c in r.cells] for r in list(t.rows)[:12]]}
    if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type):
        try: d['crop']=[round(x,2) for x in (sh.crop_left,sh.crop_right,sh.crop_top,sh.crop_bottom)]
        except Exception: pass
        try: d['rotation']=round(sh.rotation,1)
        except Exception: pass
    try:
        f=sh.fill
        if f.type is not None:
            ft=str(f.type).split('.')[-1].split(' ')[0]; d['fill']=ft
            if ft=='SOLID':
                try: d['fill']=f"SOLID:{f.fore_color.rgb}"
                except Exception:
                    try: d['fill']=f"SOLID:theme:{f.fore_color.theme_color}"
                    except Exception: pass
    except Exception: pass
    try:
        if sh.is_placeholder: d['placeholder']=str(sh.placeholder_format.type).split('.')[-1].split(' ')[0]
    except Exception: pass
    return d

def deck_summary(path):
    prs=Presentation(str(path))
    out={'slide_width_in':round(prs.slide_width/914400,2),'slide_height_in':round(prs.slide_height/914400,2),'n_slides':len(prs.slides),
         'n_layouts':sum(len(m.slide_layouts) for m in prs.slide_masters),'n_masters':len(prs.slide_masters),'slides':[]}
    for i,s in enumerate(prs.slides,1):
        shapes=[]
        for sh in s.shapes:
            try: shapes.append(shape_summary(sh))
            except Exception as e: shapes.append({'name':sh.name,'error':str(e)})
        sd={'idx':i,'layout':s.slide_layout.name,'shapes':shapes}
        try:
            fill=s.background.fill; sd['bg_fill']=str(fill.type)
        except Exception: pass
        xml=s._element.xml
        sd['has_animation']='<p:timing' in xml; sd['has_transition']='<p:transition' in xml
        out['slides'].append(sd)
    return out

def diff_text(a,b,label_a,label_b):
    """Compact human-readable diff of two deck summaries (per slide: changed/added/removed shapes)."""
    lines=[]
    if (a['slide_width_in'],a['slide_height_in'])!=(b['slide_width_in'],b['slide_height_in']):
        lines.append(f"slide size: {label_a}={a['slide_width_in']}x{a['slide_height_in']}in  {label_b}={b['slide_width_in']}x{b['slide_height_in']}in")
    if a['n_slides']!=b['n_slides']: lines.append(f"slide count: {label_a}={a['n_slides']} {label_b}={b['n_slides']}")
    for i in range(max(a['n_slides'],b['n_slides'])):
        sa=a['slides'][i] if i<a['n_slides'] else None; sb=b['slides'][i] if i<b['n_slides'] else None
        if sa is None: lines.append(f"slide {i+1}: only in {label_b}: layout={sb['layout']} shapes={[s.get('name') for s in sb['shapes']]} texts={[s.get('text','')[:60] for s in sb['shapes'] if s.get('text')]}"); continue
        if sb is None: lines.append(f"slide {i+1}: only in {label_a}: layout={sa['layout']} texts={[s.get('text','')[:60] for s in sa['shapes'] if s.get('text')]}"); continue
        sl=[]
        if sa['layout']!=sb['layout']: sl.append(f"layout {sa['layout']!r}->{sb['layout']!r}")
        if sa.get('bg_fill')!=sb.get('bg_fill'): sl.append(f"bg {sa.get('bg_fill')}->{sb.get('bg_fill')}")
        if sa['has_animation']!=sb['has_animation']: sl.append(f"animation {sa['has_animation']}->{sb['has_animation']}")
        ka={json.dumps(s,sort_keys=True) for s in sa['shapes']}; kb={json.dumps(s,sort_keys=True) for s in sb['shapes']}
        only_a=[json.loads(x) for x in ka-kb]; only_b=[json.loads(x) for x in kb-ka]
        for s in only_a: sl.append(f"  {label_a} shape: {json.dumps(s,ensure_ascii=False)[:400]}")
        for s in only_b: sl.append(f"  {label_b} shape: {json.dumps(s,ensure_ascii=False)[:400]}")
        if sl: lines.append(f"slide {i+1}: "+"\n".join(sl))
    return "\n".join(lines) if lines else "(no structural differences detected)"

def main(slug):
    m=manifest[slug]; wd=Path(m['workdir']); deck=wd/'deck.pptx'; od=OUT/slug; od.mkdir(parents=True,exist_ok=True)
    rec={'slug':slug,'name':m['name'],'edit_type':m['edit_type'],'category':m['category']}
    st=(wd/'.started'); fi=(wd/'.finished')
    rec['wall_seconds']=int(fi.read_text())-int(st.read_text()) if st.exists() and fi.exists() else None
    rec['finished_marker']=fi.exists()
    if not deck.exists(): rec['status']='missing_output'; return rec
    rec['sha_changed']=sha(deck)!=m['original_sha256']
    rec['is_zip_pptx']=zipfile.is_zipfile(deck) and '[Content_Types].xml' in zipfile.ZipFile(deck).namelist()
    try:
        pred=deck_summary(deck); rec['python_pptx_opens']=True
    except Exception as e:
        rec['python_pptx_opens']=False; rec['python_pptx_error']=str(e)[:300]; pred=None
    pngs,err=render(deck, od/'pred'); rec['renders_with_soffice']=bool(pngs); rec['n_pred_png']=len(pngs)
    if err: rec['render_error']=err
    rec['status']='no_change' if not rec['sha_changed'] else ('invalid_pptx' if not rec['is_zip_pptx'] or not rec['python_pptx_opens'] else 'ok')
    orig=deck_summary(REPO/m['original']); gt=deck_summary(REPO/m['ground_truth'])
    json.dump({'orig':orig,'gt':gt,'pred':pred}, open(od/'summaries.json','w'), indent=0)
    if pred:
        (od/'diff_orig_vs_pred.txt').write_text(diff_text(orig,pred,'ORIGINAL','PREDICTION'))
        (od/'diff_gt_vs_pred.txt').write_text(diff_text(gt,pred,'GROUND_TRUTH','PREDICTION'))
        (od/'diff_orig_vs_gt.txt').write_text(diff_text(orig,gt,'ORIGINAL','GROUND_TRUTH'))
        rec['n_slides']={'orig':orig['n_slides'],'gt':gt['n_slides'],'pred':pred['n_slides']}
        rec['slide_size_in']={'orig':[orig['slide_width_in'],orig['slide_height_in']],'gt':[gt['slide_width_in'],gt['slide_height_in']],'pred':[pred['slide_width_in'],pred['slide_height_in']]}
        cnt=lambda d,k: sum(1 for s in d['slides'] for sh in s['shapes'] if k in sh)
        rec['n_charts']={'orig':cnt(orig,'chart'),'gt':cnt(gt,'chart'),'pred':cnt(pred,'chart')}
        rec['n_tables']={'orig':cnt(orig,'table'),'gt':cnt(gt,'table'),'pred':cnt(pred,'table')}
    json.dump(rec, open(od/'record.json','w'), indent=1)
    print(json.dumps(rec, indent=1))
    return rec

if __name__=='__main__':
    for s in sys.argv[1:]: main(s)
