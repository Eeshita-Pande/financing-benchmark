"""Create the synthetic 4-slide deck used by the pipeline smoke test (SMOKE-001).

    python make_smoke_deck.py <smoke_root>

Writes <smoke_root>/data/files/PowerPoint/synthetic.pptx. Copy
``smoke_registry/task_registry`` to <smoke_root>/task_registry next to it.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

root = Path(sys.argv[1])
out = root / "data" / "files" / "PowerPoint" / "synthetic.pptx"
out.parent.mkdir(parents=True, exist_ok=True)
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Lecture 3: Synthetic Deck"
s.placeholders[1].text = "smoke test"
for i in range(2, 5):
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = f"Slide {i} title"
    tf = sl.placeholders[1].text_frame
    tf.text = "Alpha"
    tf.add_paragraph().text = "Beta"
    tf.add_paragraph().text = "Gamma"
    tb = sl.shapes.add_textbox(Inches(8), Inches(7), Inches(1.5), Inches(0.4))
    tb.text_frame.text = f"Lec 3.{i}"
prs.save(out)
print("wrote", out)
