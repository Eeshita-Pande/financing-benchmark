#!/usr/bin/env python3
"""Build results.json (all runs incl. model variants and lenient regrade) and markdown tables for the README."""
import json, re
from pathlib import Path
from collections import defaultdict
BENCH = Path('/home/user/bench')
inv = {r['task']: r for r in json.loads((BENCH / 'inventory.json').read_text())}
order = [l.strip() for l in (BENCH / 'tools' / 'tasklist.txt').read_text().splitlines() if l.strip() and not l.startswith('#')]
runs = []
for s in sorted((BENCH / 'runs').glob('*/summary.json')):
    d = json.loads(s.read_text())
    d['run_dir'] = s.parent.name
    d['level'] = inv[d['task']]['level']; d['doc_type'] = inv[d['task']]['doc_type']
    # exact agent model from the transcript: the model that did the tool-using turns
    models = defaultdict(int)
    for line in (s.parent / 'claude-code.jsonl').read_text(errors='ignore').splitlines():
        try: r = json.loads(line)
        except Exception: continue
        if r.get('type') == 'assistant':
            models[r.get('message', {}).get('model', '?')] += 1
    d['agent_model_from_transcript'] = max(models, key=models.get) if models else None
    d['assistant_messages_by_model'] = dict(models)
    d['tests_passed'] = sum(1 for v in d['tests'].values() if v['outcome'] == 'passed')
    d['tests_total'] = len(d['tests'])
    # do the files the agent actually wrote to submission/ open? (strict file-name match is reported separately)
    import zipfile
    sub = s.parent / 'work' / 'submission'
    written = [f for f in sub.iterdir() if f.is_file()] if sub.exists() else []
    def opens(f):
        try:
            if f.suffix.lower() in ('.pptx', '.xlsx', '.docx'):
                with zipfile.ZipFile(f) as z:
                    if z.testzip() is not None: return False
                if f.suffix.lower() == '.pptx':
                    from pptx import Presentation; Presentation(str(f))
                elif f.suffix.lower() == '.xlsx':
                    from openpyxl import load_workbook; load_workbook(str(f))
                else:
                    from docx import Document; Document(str(f))
            elif f.suffix.lower() == '.pdf':
                from pypdf import PdfReader; PdfReader(str(f))
            return True
        except Exception:
            return False
    d['submission_files'] = {f.name: opens(f) for f in written}
    d['outputs_open'] = bool(written) and all(d['submission_files'].values())
    d['expected_names_present'] = all(v['exists'] for v in d['files'].values())
    runs.append(d)
def short(t):
    return t.replace('cross_doc_docops__', 'xdoc:').replace('single_doc_workflow__', 'wf:').replace('atomic__', 'atom:').replace('docops_v2_', 'v2:').replace('_seed', '')
def row(d):
    len_ = d.get('lenient', {})
    lt = f"{sum(1 for v in len_['tests'].values() if v['outcome']=='passed')}/{len(len_['tests'])}" if len_.get('applied') else '-'
    lr = len_.get('reward', '-') if len_.get('applied') else '-'
    return f"| `{short(d['task'])}` | {d['level']} | {d['mode'].replace('skill_','')} | {d['agent_model_from_transcript']} | **{d['reward']}** | {d['tests_passed']}/{d['tests_total']} | {lr} | {lt} | {d['wall_seconds']:.0f} | {d['num_turns']} | {d['n_tool_uses']} | {'yes' if d['outputs_open'] else 'NO'} | {d['transcript_verifier_mentions']} |"
hdr = "| task | level | skills | model | strict reward | tests | lenient reward | lenient tests | wall s | turns | tool calls | outputs open | verifier mentions |\n|---|---|---|---|---|---|---|---|---|---|---|---|---|"
main = [d for d in runs if d['model_requested'] == 'cli-default']
variant = [d for d in runs if d['model_requested'] != 'cli-default']
lines = ['## Main runs (CLI default model)', '', hdr]
for t in order:
    for mode in ('skill_off', 'skill_on'):
        for d in main:
            if d['task'] == t and d['mode'] == mode: lines.append(row(d))
if variant:
    lines += ['', '## Model-variant reruns of verifier-failed tasks', '', hdr]
    for t in order:
        for d in main:
            if d['task'] == t and d['mode'] == 'skill_off' and d['reward'] == 0: lines.append(row(d))
        for d in variant:
            if d['task'] == t: lines.append(row(d))
# pass rates by level
lines += ['', '## Pass rate by level (main runs)', '', '| level | skills | tasks | strict pass | lenient pass |', '|---|---|---|---|---|']
for lvl in ('L1', 'L3', 'L4'):
    for mode in ('skill_off', 'skill_on'):
        ds = [d for d in main if d['level'] == lvl and d['mode'] == mode]
        if not ds: continue
        sp = sum(d['reward'] for d in ds)
        lp = sum((d.get('lenient') or {}).get('reward', d['reward']) if (d.get('lenient') or {}).get('applied') else d['reward'] for d in ds)
        lines.append(f"| {lvl} | {mode.replace('skill_','')} | {len(ds)} | {sp}/{len(ds)} | {lp}/{len(ds)} |")
(BENCH / 'results_tables.md').write_text('\n'.join(lines) + '\n')
out = {'benchmark': 'icip-cas/DocOps', 'docops_commit': 'ccf7a75', 'date': '2026-10-13', 'agent': 'Claude Code CLI 2.1.296 (headless --print, bypassPermissions)',
       'execution': 'outside Docker; see README caveats', 'runs': runs}
(BENCH / 'results.json').write_text(json.dumps(out, indent=1))
print('\n'.join(lines))
