#!/usr/bin/env python3
"""Materialise one DocOps task into an isolated workdir, run Claude Code on it, then grade.

Usage: run_task.py TASK_NAME MODE   (MODE = off | on)
Outside Docker: every /root/... path in the task is rewritten to <workdir>/...
"""
import json, os, re, shutil, subprocess, sys, time, zipfile
from pathlib import Path

BENCH = Path('/home/user/bench')
TASKS = BENCH / 'DocOps' / 'tasks'
SKILLS = BENCH / 'DocOps' / 'skills'
VENV_BIN = BENCH / 'venv' / 'bin'
AGENT_TIMEOUT = int(os.environ.get('AGENT_TIMEOUT', '900'))   # 15 min cap
MODEL = os.environ.get('DOCOPS_MODEL', '')                     # '' = CLI default

task, mode = sys.argv[1], sys.argv[2]
grade_only = len(sys.argv) > 3 and sys.argv[3] == 'grade'
assert mode in ('off', 'on')
src = TASKS / task
suffix = f'{task}__skill_{mode}' + (f'__{MODEL}' if MODEL else '')
final = BENCH / 'runs' / suffix
run = final if grade_only else BENCH / 'live' / suffix
work = run / 'work'
cfg = run / 'claude_config'
if not grade_only:
    for d in (run, final):
        if d.exists():
            shutil.rmtree(d)
    run.mkdir(parents=True)
    work.mkdir()
    cfg.mkdir()

def rewrite(text: str) -> str:
    return text.replace('/root/', f'{work}/').replace('/tests/', f'{run}/tests/')

# 1. environment inputs: ONLY the files the task Dockerfile COPYs into /root (skills handled separately).
#    environment/task_metadata.json holds verifier expectations and is shipped only when the Dockerfile says so.
copied = []
if not grade_only:
    for line in (src / 'environment' / 'Dockerfile').read_text().splitlines():
        parts = line.split()
        if len(parts) != 3 or parts[0] != 'COPY' or parts[1].startswith('skills'):
            continue
        srcf, dst = parts[1], parts[2]
        assert dst.startswith('/root/'), line
        target = work / dst[len('/root/'):]
        target.parent.mkdir(parents=True, exist_ok=True)
        sp = src / 'environment' / srcf
        if sp.is_dir():
            shutil.copytree(sp, target)
        else:
            shutil.copy2(sp, target)
        copied.append(dst)
# 2. instruction with rewritten paths
instruction = rewrite((src / 'instruction.md').read_text())
if not grade_only:
    (run / 'instruction.md').write_text(instruction)
# 3. tests with rewritten metadata
meta_p = run / 'tests' / 'task_metadata.json'
def materialise_tests():
    (run / 'tests').mkdir(exist_ok=True)
    for p in (src / 'tests').iterdir():
        shutil.copy2(p, run / 'tests' / p.name)
        if p.suffix in ('.py', '.json'):
            (run / 'tests' / p.name).write_text(rewrite(p.read_text()))
if grade_only:
    materialise_tests()
meta = json.loads(rewrite((src / 'tests' / 'task_metadata.json').read_text()))
# 4. skills (project-level) for skill-on
if mode == 'on' and not grade_only:
    shutil.copytree(SKILLS, work / '.claude' / 'skills')
(work / 'submission').mkdir(exist_ok=True)
prev = json.loads((run / 'summary.json').read_text()) if grade_only and (run / 'summary.json').exists() else {}

# 5. run Claude Code
env = dict(os.environ)
env.pop('CLAUDE_CODE_SYNC_SKILLS', None)
env['CLAUDE_CONFIG_DIR'] = str(cfg)
env['PATH'] = f'{VENV_BIN}:' + env['PATH']   # python/python3 -> venv with office libs
env['IS_SANDBOX'] = '1'
# task.toml [environment.env] (INPUT_PATH, OUTPUT_PATH, TASK_ID, DOC_TYPE) is set inside the Harbor container; mirror it.
import tomllib
tt = tomllib.loads((src / 'task.toml').read_text())
task_env = {k: rewrite(str(v)) for k, v in tt.get('environment', {}).get('env', {}).items()}
env.update(task_env)
cmd = ['claude', '--verbose', '--output-format', 'stream-json', '--permission-mode', 'bypassPermissions',
       '--setting-sources', 'project', '--strict-mcp-config']
if mode == 'off':
    cmd.append('--disable-slash-commands')
if MODEL:
    cmd += ['--model', MODEL]
cmd += ['--print', '--', instruction]
log = run / 'claude-code.jsonl'
t0 = time.time()
timed_out = False
if grade_only:
    rc = prev.get('exit_code'); timed_out = prev.get('timed_out', False)
with open(log, 'a' if grade_only else 'w') as fh:
    try:
        if grade_only:
            raise RuntimeError('skip')
        proc = subprocess.run(cmd, cwd=work, env=env, stdout=fh, stderr=subprocess.STDOUT,
                              stdin=subprocess.DEVNULL, timeout=AGENT_TIMEOUT)
        rc = proc.returncode
    except subprocess.TimeoutExpired:
        timed_out, rc = True, -1
    except RuntimeError:
        pass
wall = prev.get('wall_seconds', 0) if grade_only else time.time() - t0

if not grade_only:
    materialise_tests()
# parse the stream for the result record
result_rec, n_tool_uses, tools = None, 0, {}
for line in log.read_text(errors='ignore').splitlines():
    try:
        rec = json.loads(line)
    except Exception:
        continue
    if rec.get('type') == 'result':
        result_rec = rec
    if rec.get('type') == 'assistant':
        for blk in rec.get('message', {}).get('content', []):
            if isinstance(blk, dict) and blk.get('type') == 'tool_use':
                n_tool_uses += 1
                tools[blk.get('name')] = tools.get(blk.get('name'), 0) + 1

verifier_hits = sum(1 for line in log.read_text(errors='ignore').splitlines()
                    if re.search(r'tests/|test_outputs|verifier_utils|verifier_expectations', line))
# 6. output files: existence + opens (zip integrity + library load)
outs = meta.get('output_paths') or {'main': meta['output_path']}
files = {}
for kind, path in outs.items():
    p = Path(path)
    info = {'path': path, 'exists': p.exists(), 'size': p.stat().st_size if p.exists() else 0, 'opens': None}
    if p.exists() and p.stat().st_size:
        try:
            ext = p.suffix.lower()
            if ext in ('.pptx', '.xlsx', '.docx'):
                with zipfile.ZipFile(p) as z:
                    bad = z.testzip()
                if ext == '.pptx':
                    from pptx import Presentation; Presentation(str(p))
                elif ext == '.xlsx':
                    from openpyxl import load_workbook; load_workbook(str(p))
                else:
                    from docx import Document; Document(str(p))
                info['opens'] = bad is None
            elif ext == '.pdf':
                from pypdf import PdfReader; PdfReader(str(p)); info['opens'] = True
            else:
                info['opens'] = True
        except Exception as e:
            info['opens'] = False
            info['open_error'] = repr(e)[:200]
    files[kind] = info

# 7. grade with the task's own verifier (pytest, -rA -v like tests/test.sh), json report
genv = dict(os.environ)
genv['TASK_METADATA_PATH'] = str(meta_p)
genv['INPUT_PATH'] = meta.get('input_path', '')
genv['OUTPUT_PATH'] = meta.get('output_path', '')
# the task.toml [verifier.env] only sets TASK_METADATA_PATH; test.sh runs pytest in the container.
report = run / 'pytest_report.json'
g = subprocess.run([str(VENV_BIN / 'python'), '-I', '-m', 'pytest', str(run / 'tests' / 'test_outputs.py'),
                    '-rA', '-v', '-p', 'no:cacheprovider', '--json-report', f'--json-report-file={report}'],
                   cwd=run / 'tests', env=genv, capture_output=True, text=True, timeout=600)
(run / 'pytest_output.txt').write_text(g.stdout + g.stderr)
tests = {}
if report.exists():
    rep = json.loads(report.read_text())
    for t in rep.get('tests', []):
        name = t['nodeid'].split('::')[-1]
        msg = ''
        if t['outcome'] != 'passed':
            cr = t.get('call') or t.get('setup') or {}
            msg = (cr.get('crash') or {}).get('message', '') or (cr.get('longrepr', '') or '')[-800:]
        tests[name] = {'outcome': t['outcome'], 'message': msg[:800]}
    reward = 1 if rep['summary'].get('passed', 0) == rep['summary'].get('total', 0) and rep['summary'].get('total', 0) > 0 else 0
else:
    reward = 0

summary = {
    'task': task, 'mode': f'skill_{mode}', 'model_requested': MODEL or 'cli-default',
    'models_used': list((result_rec or {}).get('modelUsage', {}).keys()),
    'wall_seconds': round(wall, 1), 'timed_out': timed_out, 'exit_code': rc,
    'num_turns': (result_rec or {}).get('num_turns'), 'n_tool_uses': n_tool_uses, 'tools': tools,
    'cost_usd': (result_rec or {}).get('total_cost_usd'), 'is_error': (result_rec or {}).get('is_error'),
    'agent_final_text': ((result_rec or {}).get('result') or '')[:1500],
    'files': files, 'tests': tests, 'reward': reward, 'pytest_exit': g.returncode,
    'transcript_verifier_mentions': verifier_hits,
}
summary['inputs_materialised'] = copied if copied else prev.get('inputs_materialised')
summary['task_env_exported'] = task_env if not grade_only else prev.get('task_env_exported')
(run / 'summary.json').write_text(json.dumps(summary, indent=1))
if run != final:
    final.parent.mkdir(exist_ok=True)
    shutil.move(str(run), str(final))
print(json.dumps({k: summary[k] for k in ('task', 'mode', 'wall_seconds', 'timed_out', 'num_turns', 'reward')}),
      {k: v['outcome'] for k, v in tests.items()})
