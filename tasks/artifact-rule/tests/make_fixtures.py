#!/usr/bin/env python3
"""Build hand-made grader fixtures: a correct output and a deliberately misplaced output per instance.

The misplaced output is identical to the correct one except that the new slide sits where the brief
literally allows (end of the case block / end of the appendix / end of the financing section).
Usage: python3 tests/make_fixtures.py   (writes tests/fixtures/<instance>_{correct,misplaced}.pptx)
"""
import copy
import json
import os
import sys

from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
import generate as G  # noqa: E402


def duplicate_slide(prs, src):
    """Append a structural duplicate of slide `src` (same layout, deep-copied shapes, copied notes)."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):  # drop layout-created placeholders; we copy the source's instead
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    if src.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return new


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = [x for x in ids if x.rId == prs.part.relate_to(slide.part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide")][0]
    lst.remove(el)
    lst.insert(new_index, el)


def set_run_text(shape_or_cell_tf, para_index, text):
    p = shape_or_cell_tf.paragraphs[para_index]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


def fill_card(slide, title, tiles, rows, footer, notes):
    slide.shapes.title.text_frame.paragraphs[0].runs[0].text = title
    tile_shapes = [s for s in slide.shapes if s.name.startswith("Tile ")]
    for shp, val in zip(tile_shapes, tiles):
        set_run_text(shp.text_frame, 1, val)
    tbl = [s for s in slide.shapes if s.has_table][0].table
    for i, val in enumerate(rows, 1):
        set_run_text(tbl.cell(i, 1).text_frame, 0, val)
    foot = [s for s in slide.shapes if s.name == "Source note"][0]
    set_run_text(foot.text_frame, 0, footer)
    slide.notes_slide.notes_text_frame.text = notes


def fill_schedule(slide, title, header, rows, footer, notes):
    slide.shapes.title.text_frame.paragraphs[0].runs[0].text = title
    tbl = [s for s in slide.shapes if s.has_table][0].table
    for j, h in enumerate(header):
        set_run_text(tbl.cell(0, j).text_frame, 0, h)
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            set_run_text(tbl.cell(i, j).text_frame, 0, v)
    foot = [s for s in slide.shapes if s.name == "Source note"][0]
    set_run_text(foot.text_frame, 0, footer)
    slide.notes_slide.notes_text_frame.text = notes


def update_agenda_i1(prs):
    _, ag = [(i, s) for i, s in enumerate(prs.slides) if s.shapes.title.text == "Agenda"][0]
    body = [ph for ph in ag.placeholders if ph.placeholder_format.idx != 0][0]
    tf = body.text_frame
    paras = list(tf.paragraphs)
    base_p = [p for p in paras if p.text.startswith("Case 1")][0]
    new_p = copy.deepcopy(base_p._p)
    base_p._p.addnext(new_p)
    # set text of the new paragraph
    from pptx.text.text import _Paragraph
    np_ = _Paragraph(new_p, tf)
    np_.runs[0].text = "Case 4 — Renewal without credit transfer"


def build(instance, misplaced, out_path):
    inst_dir = os.path.join(ROOT, "instances", instance)
    spec = json.load(open(os.path.join(inst_dir, "expected", "spec.json")))
    prs = Presentation(os.path.join(inst_dir, "inputs", "ic_deck_base.pptx"))
    titles = [s.shapes.title.text for s in prs.slides]
    src = prs.slides[titles.index(spec["structure_template_title"])]
    new = duplicate_slide(prs, src)
    if instance == "i1_case4":
        tiles, rows = G.card_values(G.CASES[3])
        fill_card(new, spec["new_slide_title"], tiles, rows, G.CASE_FOOTER, "Case 4 headline metrics from the Case summary tab.")
        update_agenda_i1(prs)
    elif instance == "i3_sensitivity":
        tiles, rows = G.card_values(G.SENSITIVITY)
        fill_card(new, spec["new_slide_title"], tiles, rows, "Source: model_outputs.xlsx, tab 'Sensitivity'",
                  "Sensitivity of the covenant DSCR to a 20% lower support cap.")
    else:
        rows = [[r[0], G.fmt_int(r[1]), G.fmt_int(r[2]), G.fmt_pct(r[3])] for r in G.support_rows()]
        fill_schedule(new, spec["new_slide_title"], G.SUPPORT_HEADER, rows,
                      "Source: model_outputs.xlsx, tab 'Support cap'", "Support-cap utilisation per case.")
    order = spec["literal_placement_title_order"] if misplaced else spec["expected_title_order"]
    move_slide(prs, new, order.index(spec["new_slide_title"]))
    prs.save(out_path)


def main():
    out_dir = os.path.join(HERE, "fixtures")
    os.makedirs(out_dir, exist_ok=True)
    for inst in G.BRIEFS:
        for tag, mis in (("correct", False), ("misplaced", True)):
            path = os.path.join(out_dir, f"{inst}_{tag}.pptx")
            build(inst, mis, path)
            print("wrote", path)


if __name__ == "__main__":
    main()
